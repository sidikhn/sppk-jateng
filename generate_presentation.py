import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    # Color Palette definitions
    BG_DARK = RGBColor(11, 25, 44)       # Deep Navy / Ocean
    TEAL_PRIMARY = RGBColor(13, 148, 136) # Teal / Aurora (#0D9488)
    TEAL_DARK = RGBColor(15, 118, 110)   # Deep Teal (#0F766E)
    TEAL_LIGHT = RGBColor(204, 251, 241) # Light Teal background
    CYAN_ACCENT = RGBColor(6, 182, 212)  # Cyan
    CARD_BG = RGBColor(248, 250, 252)    # Slate 50
    CARD_BORDER = RGBColor(226, 232, 240)# Slate 200
    TEXT_MAIN = RGBColor(15, 23, 42)     # Slate 900
    TEXT_MUTED = RGBColor(71, 85, 105)   # Slate 600
    TEXT_WHITE = RGBColor(255, 255, 255)
    GOLD_ACCENT = RGBColor(217, 119, 6)  # Amber 600
    RED_ACCENT = RGBColor(225, 29, 72)   # Rose 600
    GREEN_ACCENT = RGBColor(16, 185, 129)# Emerald 500

    def add_header(slide, tag_text, title_text, subtitle_text=""):
        # Header banner area
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = tag_text.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = TEAL_PRIMARY
        p0.space_after = Pt(2)

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_MAIN

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(11)
            p2.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # =========================================================================
    # SLIDE 1: COVER SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = BG_DARK
    bg1.line.fill.background()

    # Decorative accent stripe
    stripe = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.12), Inches(4.0))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = TEAL_PRIMARY
    stripe.line.fill.background()

    tbox1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.7), Inches(11.2), Inches(4.5))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "SISTEM PENDUKUNG KEPUTUSAN (SPPK)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p.space_after = Pt(8)

    p = tf1.add_paragraph()
    p.text = "Penentuan Prioritas Intervensi Penanganan Stunting\ndi Provinsi Jawa Tengah"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_after = Pt(12)

    p = tf1.add_paragraph()
    p.text = "Integrasi Metode Analytic Hierarchy Process (AHP) & TOPSIS Berbasis Web"
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(148, 163, 184) # Slate 400
    p.space_after = Pt(28)

    p = tf1.add_paragraph()
    p.text = "Disusun untuk Presentasi Akademik & Ujian Sistem Pendukung Keputusan\nProgram Pascasarjana Universitas Gadjah Mada (UGM)"
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(203, 213, 225) # Slate 300

    notes1 = s1.notes_slide.notes_text_frame
    notes1.text = (
        "SKRIP BICARA (COVER):\n"
        "Selamat pagi/siang Bapak/Ibu dosen penguji dan pembimbing. "
        "Terima kasih atas kesempatan yang diberikan. Pada kesempatan kali ini, saya akan mempresentasikan "
        "penelitian dan rancang bangun Sistem Pendukung Keputusan Penentuan Prioritas Intervensi Penanganan Stunting "
        "di Wilayah Jawa Tengah dengan mengintegrasikan metode AHP (Analytic Hierarchy Process) dan TOPSIS berbasis web. "
        "Fokus utama sistem ini adalah menghadirkan rekomendasi alokasi intervensi stunting yang objektif, multi-kriteria, "
        "dan dapat dipertanggungjawabkan secara matematis."
    )

    # =========================================================================
    # SLIDE 2: LATAR BELAKANG & URGENSI MASALAH
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Pendahuluan & Konteks", "Latar Belakang & Urgensi Masalah", "Mengapa penentuan prioritas stunting di Jawa Tengah krusial?")

    # 3 Column Cards
    col_w = Inches(3.64)
    gap = Inches(0.39)
    top_pos = Inches(1.7)
    h_pos = Inches(4.9)

    # Card 1: Isu Stunting
    c1 = add_card(s2, Inches(0.8), top_pos, col_w, h_pos)
    tb = s2.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), col_w - Inches(0.4), h_pos - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. Isu Strategis Nasional"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RED_ACCENT
    p.space_after = Pt(10)
    bullet_items = [
        "Target Nasional: Penurunan stunting menuju angka di bawah 14% (Perpres No. 72/2021).",
        "Stunting bukan hanya masalah tinggi badan, melainkan gangguan perkembangan kognitif permanen dan produktivitas masa depan.",
        "Dibutuhkan akselerasi intervensi yang terarah, cepat, dan presisi tinggi."
    ]
    for b in bullet_items:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    # Card 2: Disparitas 35 Daerah
    c2 = add_card(s2, Inches(0.8) + col_w + gap, top_pos, col_w, h_pos)
    tb = s2.shapes.add_textbox(Inches(1.0) + col_w + gap, top_pos + Inches(0.2), col_w - Inches(0.4), h_pos - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. Disparitas 35 Wilayah"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT
    p.space_after = Pt(10)
    bullet_items2 = [
        "Jawa Tengah memiliki 35 Kabupaten/Kota dengan kondisi geografis dan fiskal yang sangat heterogen.",
        "Variasi indikator sangat tinggi: prevalensi stunting berkisar 8.9% hingga 23.5%, dengan disparitas kemiskinan dan nakes yang lebar.",
        "Alokasi anggaran bantuan penanganan seringkali terbatas dan tidak merata jika tidak diatur kuadran prioritasnya."
    ]
    for b in bullet_items2:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    # Card 3: Problem Pengambilan Keputusan
    c3 = add_card(s2, Inches(0.8) + (col_w + gap)*2, top_pos, col_w, h_pos)
    tb = s2.shapes.add_textbox(Inches(1.0) + (col_w + gap)*2, top_pos + Inches(0.2), col_w - Inches(0.4), h_pos - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "3. Gap Keputusan Konvensional"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY
    p.space_after = Pt(10)
    bullet_items3 = [
        "Keputusan penentuan lokasi fokus (lokus) seringkali berbasis intuisi, parsial, atau lobi birokrasi semata.",
        "Sulit memadukan kriteria yang sifatnya multi-dimensi (kesehatan balita, sanitasi, air bersih, daya dukung nakes).",
        "Solusi: SPPK berbasis Multi-Criteria Decision Making (MCDM) untuk menjamin transparansi & akuntabilitas."
    ]
    for b in bullet_items3:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    notes2 = s2.notes_slide.notes_text_frame
    notes2.text = (
        "SKRIP BICARA (LATAR BELAKANG):\n"
        "Latar belakang penelitian ini berpijak pada 3 poin krusial. Pertama, stunting adalah agenda nasional sesuai Perpres 72/2021. "
        "Kedua, di Jawa Tengah terdapat 35 Kabupaten/Kota dengan disparitas tajam. Prevalensi stunting terendah ada di angka 8.9% di Kota Semarang, "
        "namun di daerah Brebes mencapai 23.5%. Ketiga, keterbatasan anggaran menuntut pemerintah menetapkan prioritas intervensi. "
        "Namun selama ini, penentuan prioritas kerap dilakukan secara parsial hanya melihat 1 indikator saja atau dipengaruhi subjektivitas birokrasi. "
        "Oleh karena itu, diperlukan Sistem Pendukung Keputusan yang objektif, saintifik, dan transparan."
    )

    # =========================================================================
    # SLIDE 3: RUMUSAN MASALAH & TUJUAN PENELITIAN
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Fokus Penelitian", "Rumusan Masalah & Tujuan Sistem", "Landasan perancangan model pengambilan keputusan multi-kriteria")

    # 2 Big Cards: Rumusan Masalah vs Tujuan
    card_w2 = Inches(5.6)
    c_left = add_card(s3, Inches(0.8), Inches(1.7), card_w2, Inches(5.0))
    tb_l = s3.shapes.add_textbox(Inches(1.1), Inches(1.9), card_w2 - Inches(0.6), Inches(4.6))
    tfl = tb_l.text_frame
    tfl.word_wrap = True
    p = tfl.paragraphs[0]
    p.text = "RUMUSAN MASALAH"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RED_ACCENT
    p.space_after = Pt(14)

    rms = [
        "1. Bagaimana menyusun pembobotan antar-indikator stunting secara terstruktur dan teruji konsistensi logisnya?",
        "2. Bagaimana melakukan pemeringkatan secara adil dan simultan terhadap 35 Kabupaten/Kota dengan indikator yang saling berlawanan (benefit vs cost)?",
        "3. Bagaimana mengimplementasikan model matematis tersebut ke dalam platform web interaktif yang mudah digunakan oleh dinas kesehatan dan pengambil kebijakan?"
    ]
    for r in rms:
        p = tfl.add_paragraph()
        p.text = r
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(14)

    c_right = add_card(s3, Inches(6.8), Inches(1.7), card_w2, Inches(5.0))
    tb_r = s3.shapes.add_textbox(Inches(7.1), Inches(1.9), card_w2 - Inches(0.6), Inches(4.6))
    tfr = tb_r.text_frame
    tfr.word_wrap = True
    p = tfr.paragraphs[0]
    p.text = "TUJUAN PENELITIAN / SISTEM"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY
    p.space_after = Pt(14)

    tjs = [
        "1. Menerapkan metode AHP untuk membobotkan kriteria kesehatan, demografi, dan lingkungan dengan validasi Rasio Konsistensi (CR < 0.1).",
        "2. Mengintegrasikan metode TOPSIS untuk merangking 35 alternatif Kabupaten/Kota berdasarkan jarak euclidean ke solusi ideal.",
        "3. Mengembangkan sistem web SPPK interaktif lengkap dengan visualisasi radar komparasi, simulasi matriks pakar, serta ekspor laporan PDF dan Excel."
    ]
    for t in tjs:
        p = tfr.add_paragraph()
        p.text = t
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(14)

    notes3 = s3.notes_slide.notes_text_frame
    notes3.text = (
        "SKRIP BICARA (RUMUSAN & TUJUAN):\n"
        "Bapak/Ibu dosen, rumusan masalah kami bagi menjadi 3 pilar: pembobotan kriteria, pemeringkatan 35 wilayah, dan implementasi sistem web. "
        "Tujuan akhirnya adalah tidak hanya menyelesaikan formulasi matematis di atas kertas, tetapi juga menyediakan sistem operasional "
        "yang dapat langsung dipakai dinas kesehatan untuk simulasi bobot, melihat detail perhitungan matriks, hingga mencetak laporan resmi."
    )

    # =========================================================================
    # SLIDE 4: JUSTIFIKASI METODOLOGIS: HYBRID AHP-TOPSIS
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Kajian Metodologi", "Mengapa Kombinasi AHP & TOPSIS?", "Justifikasi ilmiah penggabungan kedua metode Multi-Criteria Decision Making (MCDM)")

    # 3 Cards: AHP Role, Why Not Pure AHP?, TOPSIS Role
    col_w4 = Inches(3.64)
    c1 = add_card(s4, Inches(0.8), Inches(1.7), col_w4, Inches(4.0))
    tb1 = s4.shapes.add_textbox(Inches(1.0), Inches(1.9), col_w4 - Inches(0.4), Inches(3.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Peran AHP (Weighting)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY
    p.space_after = Pt(8)
    ahp_points = [
        "Unggul dalam dekomposisi hierarki kriteria yang kompleks.",
        "Menilai kepentingan kriteria melalui Pairwise Comparison (Skala Saaty 1-9).",
        "Memiliki mekanisme Uji Konsistensi (CR < 0.1) untuk memastikan rasionalitas pakar."
    ]
    for pt in ahp_points:
        p = tf1.add_paragraph()
        p.text = "✓ " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    c2 = add_card(s4, Inches(0.8) + col_w4 + gap, Inches(1.7), col_w4, Inches(4.0))
    tb2 = s4.shapes.add_textbox(Inches(1.0) + col_w4 + gap, Inches(1.9), col_w4 - Inches(0.4), Inches(3.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Keterbatasan Pure AHP"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RED_ACCENT
    p.space_after = Pt(8)
    limit_points = [
        "Jika AHP murni digunakan merangking 35 Kabupaten/Kota:",
        "Dibutuhkan n(n-1)/2 perbandingan: 35 x 34 / 2 = 595 perbandingan per kriteria!",
        "Total 6 kriteria = 3.570 kali input perbandingan manual.",
        "Sangat rentan inkonsistensi logis dan cognitive overload bagi penilai."
    ]
    for pt in limit_points:
        p = tf2.add_paragraph()
        p.text = "✗ " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    c3 = add_card(s4, Inches(0.8) + (col_w4 + gap)*2, Inches(1.7), col_w4, Inches(4.0))
    tb3 = s4.shapes.add_textbox(Inches(1.0) + (col_w4 + gap)*2, Inches(1.9), col_w4 - Inches(0.4), Inches(3.6))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "Peran TOPSIS (Ranking)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p.space_after = Pt(8)
    topsis_points = [
        "Sangat efisien untuk perangkingan alternatif berjumlah besar (35 Kab/Kota).",
        "Menghitung jarak euclidean ke Solusi Ideal Positif (A+) & Ideal Negatif (A-).",
        "Menghilangkan kebutuhan pairwise ratusan alternatif, cukup memasukkan nilai data riil."
    ]
    for pt in topsis_points:
        p = tf3.add_paragraph()
        p.text = "★ " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    # Bottom Synergy Banner
    syn_card = add_card(s4, Inches(0.8), Inches(5.9), Inches(11.73), Inches(0.9), bg_color=TEAL_LIGHT, border_color=TEAL_PRIMARY)
    tb_syn = s4.shapes.add_textbox(Inches(1.0), Inches(5.95), Inches(11.3), Inches(0.8))
    tfs = tb_syn.text_frame
    tfs.word_wrap = True
    p = tfs.paragraphs[0]
    p.text = "KESEIMPULAN HYBRID: AHP bertugas sebagai penentu bobot objektif pakar (Wj), dan TOPSIS bertugas sebagai mesin perankingan matematis 35 daerah secara instan dan bebas bias."
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEAL_DARK

    notes4 = s4.notes_slide.notes_text_frame
    notes4.text = (
        "SKRIP BICARA (JUSTIFIKASI METODE):\n"
        "Ini adalah pertanyaan yang sering ditanyakan dosen penguji: 'Kenapa harus gabungan AHP dan TOPSIS? Kenapa tidak AHP saja?'. "
        "Jawabannya sangat fundamental secara komputasi. AHP sangat kuat untuk pembobotan kriteria berpasangan karena ada uji konsistensi CR. "
        "Tetapi jika kita menggunakan AHP untuk merangking 35 Kabupaten/Kota, maka kita harus membandingkan 35 alternatif berpasangan. "
        "Rumusnya n kali n minus 1 dibagi 2, yaitu 595 perbandingan per kriteria. Dengan 6 kriteria, pakar harus mengisi 3.570 sel matriks! "
        "Secara psikologis dan komputasi, itu hampir mustahil konsisten. "
        "Oleh karena itu, kombinasi hybrid adalah pilihan paling tepat: AHP hanya membobotkan 6 kriteria (15 perbandingan), "
        "lalu bobot tersebut diumpankan ke TOPSIS untuk merangking 35 daerah menggunakan konsep jarak solusi ideal."
    )

    # =========================================================================
    # SLIDE 5: KRITERIA & DEFINISI OPERASIONAL
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Parameter Penilaian", "Kriteria & Definisi Operasional", "6 Indikator penentu prioritas intervensi penanganan stunting di Jawa Tengah")

    # Table of Criteria
    rows = 7
    cols = 5
    table_shape = s5.shapes.add_table(rows, cols, Inches(0.8), Inches(1.7), Inches(11.73), Inches(4.9))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(1.1)  # Kode
    tbl.columns[1].width = Inches(3.2)  # Nama Kriteria
    tbl.columns[2].width = Inches(1.4)  # Sifat
    tbl.columns[3].width = Inches(1.4)  # Bobot AHP
    tbl.columns[4].width = Inches(4.63) # Rasional & Keterangan

    headers = ["Kode", "Nama Indikator / Kriteria", "Sifat", "Bobot AHP", "Rasional Penentuan & Logika SPPK"]
    for c_idx, h in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = TEAL_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.alignment = PP_ALIGN.CENTER if c_idx in [0, 2, 3] else PP_ALIGN.LEFT

    crit_data = [
        ("C1", "Prevalensi Stunting (%)", "Benefit", "37.46%", "Semakin tinggi persentase balita stunting di suatu daerah, semakin darurat intervensi."),
        ("C2", "Jumlah Balita Stunting (Jiwa)", "Benefit", "24.24%", "Semakin banyak jumlah absolut balita stunting, semakin masif beban beban populasi terdampak."),
        ("C3", "Persentase Penduduk Miskin (%)", "Benefit", "15.08%", "Semakin tinggi kemiskinan, semakin rentan ketahanan pangan dan kemampuan asupan gizi anak."),
        ("C4", "Akses Sanitasi Layak (%)", "Cost", "8.89%", "Semakin rendah sanitasi layak, semakin tinggi risiko infeksi berulang pemicu gagal tumbuh."),
        ("C5", "Akses Air Minum Layak (%)", "Cost", "8.89%", "Semakin rendah akses air minum sehat, semakin rentan sanitasi dan higienitas keluarga balita."),
        ("C6", "Jumlah Tenaga Kesehatan (Jiwa)", "Cost", "5.44%", "Semakin sedikit ketersediaan tenaga kesehatan, semakin butuh dropping bantuan intervensi medis."),
    ]

    for r_idx, row in enumerate(crit_data, start=1):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10)
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = TEAL_PRIMARY
                p.alignment = PP_ALIGN.CENTER
            elif c_idx == 2:
                p.font.bold = True
                p.font.color.rgb = GREEN_ACCENT if val == "Benefit" else RED_ACCENT
                p.alignment = PP_ALIGN.CENTER
            elif c_idx == 3:
                p.font.bold = True
                p.font.color.rgb = TEXT_MAIN
                p.alignment = PP_ALIGN.CENTER
            else:
                p.font.color.rgb = TEXT_MAIN

    notes5 = s5.notes_slide.notes_text_frame
    notes5.text = (
        "SKRIP BICARA (KRITERIA & BENEFIT/COST):\n"
        "Bapak/Ibu dosen, kami menetapkan 6 kriteria yang mencakup indikator spesifik dan sensitif stunting. "
        "Poin penting yang sering diuji adalah klasifikasi Benefit versus Cost: "
        "C1, C2, C3 diklasifikasikan sebagai BENEFIT, karena dalam konteks SPPK penetapan prioritas penanganan, "
        "semakin tinggi prevalensi stunting, jumlah balita terdampak, dan angka kemiskinan, maka daerah tersebut semakin layak mendapatkan intervensi prioritas. "
        "Sebaliknya, C4 (Sanitasi), C5 (Air Minum), dan C6 (Tenaga Kesehatan) diklasifikasikan sebagai COST, "
        "karena jika akses air, sanitasi, dan jumlah nakes di suatu daerah sangat rendah, daerah tersebut justru sangat darurat dan harus diprioritaskan."
    )

    # =========================================================================
    # SLIDE 6: TAHAP 1 - AHP: PEMBOBOTAN & VALIDASI KONSISTENSI
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Tahap 1: AHP", "Penentuan Bobot Kriteria & Uji Konsistensi Logis", "Perbandingan berpasangan (Pairwise Comparison) skala Saaty 1-9")

    # Left Card: Formula & Consistency KPI
    c_ahp_l = add_card(s6, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.0))
    tb_al = s6.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(5.2), Inches(4.7))
    tfal = tb_al.text_frame
    tfal.word_wrap = True

    p = tfal.paragraphs[0]
    p.text = "PENGUJIAN KONSISTENSI MATRIKS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY
    p.space_after = Pt(8)

    kpis = [
        ("Ukuran Matriks (n)", "6 x 6 Kriteria"),
        ("Lambda Max (λ max)", "6.0660"),
        ("Consistency Index (CI)", "0.0132  →  (λ max - n) / (n - 1)"),
        ("Random Index Saaty (RI)", "1.24  (untuk matriks n = 6)"),
        ("Consistency Ratio (CR)", "0.0106 (1.06%)"),
        ("Status Konsistensi", "SANGAT KONSISTEN (CR < 0.10 / 10%)")
    ]
    for label, val in kpis:
        p = tfal.add_paragraph()
        p.text = f"{label}: "
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        run = p.add_run()
        run.text = val
        run.font.bold = (label == "Status Konsistensi")
        run.font.color.rgb = GREEN_ACCENT if label == "Status Konsistensi" else TEXT_MUTED
        p.space_after = Pt(6)

    p = tfal.add_paragraph()
    p.text = "Interpretasi Konsistensi:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_MAIN
    p.space_after = Pt(4)

    p = tfal.add_paragraph()
    p.text = "Nilai CR sebesar 0.0106 berada jauh di bawah ambang batas toleransi Saaty (0.10). Hal ini membuktikan bahwa perbandingan antar kriteria oleh pakar bebas dari kontradiksi logis (transitive consistency terpenuhi)."
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_MUTED

    # Right Card: Weights Distribution Bar Visual
    c_ahp_r = add_card(s6, Inches(6.8), Inches(1.7), Inches(5.73), Inches(5.0))
    tb_ar = s6.shapes.add_textbox(Inches(7.0), Inches(1.85), Inches(5.3), Inches(4.7))
    tfar = tb_ar.text_frame
    tfar.word_wrap = True

    p = tfar.paragraphs[0]
    p.text = "DISTRIBUSI BOBOT HASIL EIGEN VECTOR"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY
    p.space_after = Pt(10)

    weights_bars = [
        ("C1 - Prevalensi Stunting", "37.46%", "Prioritas Utama (Determinan Klinis Primer)"),
        ("C2 - Jumlah Balita Stunting", "24.24%", "Prioritas Tinggi (Skala Populasi Terdampak)"),
        ("C3 - Persentase Kemiskinan", "15.08%", "Determinan Sosial-Ekonomi"),
        ("C4 - Akses Sanitasi Layak", "8.89%", "Faktor Lingkungan Fisik & Higienitas"),
        ("C5 - Akses Air Minum Layak", "8.89%", "Faktor Lingkungan Fisik & Higienitas"),
        ("C6 - Jumlah Tenaga Kesehatan", "5.44%", "Kapasitas Pelayanan Kesehatan"),
    ]
    for crit, pct, note in weights_bars:
        p = tfar.add_paragraph()
        p.text = f"{crit}: {pct}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(1)

        p_sub = tfar.add_paragraph()
        p_sub.text = f"→ {note}"
        p_sub.font.size = Pt(9.5)
        p_sub.font.color.rgb = TEXT_MUTED
        p_sub.space_after = Pt(6)

    notes6 = s6.notes_slide.notes_text_frame
    notes6.text = (
        "SKRIP BICARA (TAHAP AHP & KONSISTENSI):\n"
        "Pada tahap pembobotan AHP, kita menggunakan matriks 6x6 skala Saaty. "
        "Hasil normalisasi matriks dan vektor eigen menghasilkan bobot terbesar pada C1 (Prevalensi Stunting) sebesar 37.46%, "
        "disusul C2 (Balita Stunting) 24.24%, dan C3 (Kemiskinan) 15.08%. "
        "Yang terpenting untuk dipaparkan kepada dosen adalah Uji Konsistensi: "
        "Lambda Max kami peroleh 6.0660, CI sebesar 0.0132, dan dengan RI untuk n=6 adalah 1.24, "
        "diperoleh Rasio Konsistensi (CR) sebesar 0.0106 atau 1.06%. "
        "Karena CR jauh lebih kecil dari 0.10 (10%), maka struktur pertimbangan pakar terbukti valid dan konsisten secara logis."
    )

    # =========================================================================
    # SLIDE 7: TAHAP 2 - TOPSIS: TAHAPAN ALGORITMA PERANGKINGAN
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Tahap 2: TOPSIS", "Alur Algoritma Perangkingan TOPSIS", "6 Langkah matematis dari Matriks Keputusan hingga Skor Preferensi Relatif (Vi)")

    # 6 Steps in 2 Rows x 3 Cols
    sw = Inches(3.64)
    sh = Inches(2.3)
    sgap_x = Inches(0.4)
    sgap_y = Inches(0.3)

    steps_data = [
        ("1. Matriks Keputusan (X)", "Membentuk matriks X berukuran 35 x 6 dari nilai riil indikator tahun evaluasi (Xij).", TEAL_PRIMARY),
        ("2. Normalisasi Matriks (R)", "rij = xij / √(Σ xij²)\nMembagi setiap nilai elemen dengan akar jumlah kuadrat seluruh alternatif per kriteria.", TEAL_PRIMARY),
        ("3. Matriks Terbobot (Y)", "yij = wj · rij\nMengalikan matriks normalisasi R dengan bobot wj yang diperoleh dari metode AHP.", TEAL_PRIMARY),
        ("4. Solusi Ideal (A+ & A-)", "A+ = {max(yj) benefit, min(yj) cost}\nA- = {min(yj) benefit, max(yj) cost}\nMenentukan batas performa terbaik dan terburuk.", GOLD_ACCENT),
        ("5. Jarak Euclidean (D+ & D-)", "Di+ = √(Σ(yij - yj+)²)\nDi- = √(Σ(yij - yj-)²)\nMenghitung jarak geometris tiap daerah ke A+ dan A-.", GOLD_ACCENT),
        ("6. Preferensi Relatif (Vi)", "Vi = Di- / (Di+ + Di-)\nNilai preferensi 0 ≤ Vi ≤ 1. Nilai Vi tertinggi menduduki peringkat prioritas nomor 1.", RED_ACCENT),
    ]

    for idx, (stitle, sdesc, scolor) in enumerate(steps_data):
        row_i = idx // 3
        col_i = idx % 3
        pos_x = Inches(0.8) + col_i * (sw + sgap_x)
        pos_y = Inches(1.7) + row_i * (sh + sgap_y)

        c = add_card(s7, pos_x, pos_y, sw, sh)
        tb = s7.shapes.add_textbox(pos_x + Inches(0.15), pos_y + Inches(0.15), sw - Inches(0.3), sh - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = scolor
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = sdesc
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_MUTED

    notes7 = s7.notes_slide.notes_text_frame
    notes7.text = (
        "SKRIP BICARA (ALGORITMA TOPSIS):\n"
        "Setelah bobot AHP diperoleh, kita masuk ke algoritma TOPSIS yang terdiri dari 6 tahapan komputasi. "
        "Pertama, pembentukan matriks keputusan X 35 daerah kali 6 indikator. "
        "Kedua, normalisasi vektor R agar seluruh satuan yang berbeda (persen, jiwa, fasilitas) menjadi seragam tak berdimensi. "
        "Ketiga, pengalian dengan bobot AHP menghasilkan matriks terbobot Y. "
        "Keempat, penetapan Solusi Ideal Positif A+ dan Negatif A- yang disesuaikan dengan jenis benefit/cost. "
        "Kelima, kalkulasi jarak Euclidean D+ dan D-. "
        "Dan keenam, menghitung nilai preferensi relatif Vi = D minus per (D plus ditambah D minus). "
        "Daerah dengan skor Vi paling tinggi mengindikasikan kedekatan paling rapat dengan kondisi darurat stunting sehingga menjadi Prioritas 1."
    )

    # =========================================================================
    # SLIDE 8: HASIL ANALISIS & TEMUAN PERANGKINGAN 2025
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Hasil & Temuan", "Hasil Perangkingan Prioritas Stunting Jawa Tengah", "Evaluasi empiris 35 Kabupaten/Kota berbasis dataset indikator tahun 2025")

    # Table of Top 8 Districts + Category Breakdown
    tbl_shape8 = s8.shapes.add_table(9, 6, Inches(0.8), Inches(1.7), Inches(7.5), Inches(5.0))
    t8 = tbl_shape8.table
    t8.columns[0].width = Inches(0.8) # Rank
    t8.columns[1].width = Inches(2.2) # Nama
    t8.columns[2].width = Inches(1.1) # D+
    t8.columns[3].width = Inches(1.1) # D-
    t8.columns[4].width = Inches(1.1) # Vi Skor
    t8.columns[5].width = Inches(1.2) # Kategori

    headers8 = ["Rank", "Kabupaten/Kota", "Jarak D+", "Jarak D-", "Skor (Vi)", "Kategori"]
    for c_idx, h in enumerate(headers8):
        cell = t8.cell(0, c_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = TEAL_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.alignment = PP_ALIGN.CENTER if c_idx != 1 else PP_ALIGN.LEFT

    top_results = [
        ("1", "Kab. Brebes", "0.0094", "0.1531", "0.9422", "Sangat Tinggi"),
        ("2", "Kab. Pemalang", "0.0336", "0.1388", "0.8055", "Sangat Tinggi"),
        ("3", "Kab. Kebumen", "0.0450", "0.1274", "0.7392", "Sangat Tinggi"),
        ("4", "Kab. Tegal", "0.0468", "0.1287", "0.7336", "Sangat Tinggi"),
        ("5", "Kab. Wonosobo", "0.0526", "0.1213", "0.6974", "Sangat Tinggi"),
        ("6", "Kab. Grobogan", "0.0583", "0.1160", "0.6655", "Sangat Tinggi"),
        ("7", "Kab. Banjarnegara", "0.0589", "0.1144", "0.6600", "Sangat Tinggi"),
        ("8", "Kab. Cilacap", "0.0674", "0.1068", "0.6130", "Sangat Tinggi"),
    ]

    for r_idx, row in enumerate(top_results, start=1):
        for c_idx, val in enumerate(row):
            cell = t8.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(9.5)
            if c_idx in [0, 2, 3, 4, 5]:
                p.alignment = PP_ALIGN.CENTER
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = GOLD_ACCENT if r_idx <= 3 else TEXT_MAIN
            elif c_idx == 4:
                p.font.bold = True
                p.font.color.rgb = TEAL_PRIMARY
            elif c_idx == 5:
                p.font.bold = True
                p.font.color.rgb = RED_ACCENT

    # Right Card: Analisis Temuan Penting
    c_find = add_card(s8, Inches(8.6), Inches(1.7), Inches(3.93), Inches(5.0))
    tb_find = s8.shapes.add_textbox(Inches(8.8), Inches(1.85), Inches(3.53), Inches(4.7))
    tff = tb_find.text_frame
    tff.word_wrap = True

    p = tff.paragraphs[0]
    p.text = "INSIGHT HASIL PERANGKINGAN"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY
    p.space_after = Pt(8)

    insights = [
        "1. Klaster Pantura Barat & Pegunungan Tengah:\nKab. Brebes (0.9422), Pemalang (0.8055), dan Kebumen (0.7392) menduduki 3 peringkat teratas dengan kategori Sangat Tinggi karena kombinasi prevalensi tinggi (>21%), populasi balita stunting besar, dan sanitasi terbatas.",
        "2. Klaster Perkotaan Berisiko Rendah:\nKota Surakarta (0.1395), Kota Magelang (0.1526), dan Kota Semarang (0.1819) menempati peringkat terbawah (Prioritas Rendah) berkat cakupan sanitasi & air minum >95% dan rasio nakes yang sangat memadai.",
        "3. Implikasi Anggaran:\nDana DAK fisik kesehatan dan bantuan pangan intervensi gizi wajib dialokasikan prioritas ke Top 8 wilayah Sangat Tinggi."
    ]
    for ins in insights:
        p = tff.add_paragraph()
        p.text = ins
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(8)

    notes8 = s8.notes_slide.notes_text_frame
    notes8.text = (
        "SKRIP BICARA (HASIL PERANGKINGAN):\n"
        "Berdasarkan komputasi 35 Kabupaten/Kota di Jawa Tengah tahun 2025: "
        "Kabupaten Brebes menempati Ranking 1 Prioritas dengan skor preferensi 0.9422, "
        "diikuti Pemalang (0.8055), Kebumen (0.7392), Tegal (0.7336), dan Wonosobo (0.6974). "
        "Temuan ini sangat sejalan dengan realitas lapangan: wilayah Pantura Barat dan Pegunungan Tengah Selatan "
        "memiliki beban ganda berupa kemiskinan yang tinggi dan tantangan sanitasi. "
        "Sebaliknya, kota-kota otonom seperti Surakarta dan Semarang berada di kategori Rendah karena infrastruktur dasar dan nakes yang telah mapan."
    )

    # =========================================================================
    # SLIDE 9: FITUR & IMPLEMENTASI SISTEM WEB
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Implementasi Perangkat Lunak", "Arsitektur Sistem & Fitur Aplikasi Web", "Desain modular modern berbasis React, TypeScript, Tailwind CSS, & Recharts")

    # 4 Feature Cards (2x2 Grid)
    fw = Inches(5.6)
    fh = Inches(2.3)
    fgap_x = Inches(0.53)
    fgap_y = Inches(0.35)

    feats = [
        ("1. Dashboard Visual & Radar Multi-Dimensi",
         "• Visualisasi Recharts interaktif: Bar Chart Top 10 Prioritas dan Area Chart sebaran 35 daerah.\n• Radar Chart komparatif: Membandingkan secara langsung 2 Kabupaten/Kota di 6 dimensi indikator untuk mendeteksi kelemahan spesifik daerah.",
         TEAL_PRIMARY),
        ("2. Modul Simulasi AHP Interaktif",
         "• Form matriks Saaty interaktif dengan pengisian dinamis dan update sel resiprokal otomatis (1/x).\n• Validasi real-time nilai Lambda Max, CI, dan CR dengan indikator status badge Konsisten / Tidak Konsisten.",
         CYAN_ACCENT),
        ("3. Transparansi Step-by-Step TOPSIS",
         "• Tidak seperti sistem 'black box', modul ini menampilkan 7 tab tahapan kalkulasi transparan:\n  - Matriks Keputusan (X), Matriks Ternormalisasi (R), Matriks Terbobot (Y), Nilai Ideal (A+/A-), hingga Jarak D+/D-.",
         GOLD_ACCENT),
        ("4. Pelaporan Resmi & Master Data Dinamis",
         "• Master data CRUD Kabupaten/Kota, Kriteria, dan Dataset multi-tahun (2023, 2024, 2025).\n• Fitur Export PDF Resmi (dengan format surat dinas) dan Export Excel Spreadsheet dengan satu klik.",
         GREEN_ACCENT),
    ]

    for idx, (ftitle, fdesc, fcol) in enumerate(feats):
        row_i = idx // 2
        col_i = idx % 2
        pos_x = Inches(0.8) + col_i * (fw + fgap_x)
        pos_y = Inches(1.7) + row_i * (fh + fgap_y)

        c = add_card(s9, pos_x, pos_y, fw, fh)
        tb = s9.shapes.add_textbox(pos_x + Inches(0.2), pos_y + Inches(0.15), fw - Inches(0.4), fh - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = ftitle
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = fcol
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = fdesc
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MAIN

    notes9 = s9.notes_slide.notes_text_frame
    notes9.text = (
        "SKRIP BICARA (FITUR APLIKASI):\n"
        "Nilai tambah dari sistem yang kami bangun adalah transparansi dan interaktivitas: "
        "Pertama, terdapat Dashboard dengan visualisasi Radar Chart yang memungkinkan pengambil keputusan membandingkan dua kabupaten secara head-to-head. "
        "Kedua, modul AHP yang interaktif, di mana pakar dapat mengubah perbandingan matriks dan langsung melihat apakah bobotnya masih konsisten (CR < 0.1). "
        "Ketiga, transparansi penuh TOPSIS: dosen dan user dapat memeriksa tahapan matriks dari normalisasi R, Y, hingga A+ dan A-, "
        "sehingga sistem ini memenuhi prinsip Explainable AI / Transparent DSS, bukan black-box. "
        "Dan keempat, modul reporting instan ke PDF resmi dan file Excel."
    )

    # =========================================================================
    # SLIDE 10: ANALISIS SENSITIVITAS & KEKUATAN MODEL
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Validasi Model", "Analisis Sensitivitas & Stabilitas Keputusan", "Menguji ketahanan hasil perankingan terhadap fluktuasi bobot kriteria")

    # 3 Column Cards
    c1 = add_card(s10, Inches(0.8), Inches(1.7), col_w, Inches(5.0))
    tb1 = s10.shapes.add_textbox(Inches(1.0), Inches(1.9), col_w - Inches(0.4), Inches(4.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "1. Uji Sensitivitas Bobot"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY
    p.space_after = Pt(8)
    p = tf1.add_paragraph()
    p.text = "• Saat bobot kriteria Prevalensi Stunting (C1) divariasikan ±15%, peringkat Top 3 (Brebes, Pemalang, Kebumen) tetap kokoh di posisi puncak prioritas."
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_MAIN
    p.space_after = Pt(6)
    p = tf1.add_paragraph()
    p.text = "• Hal ini membuktikan model memiliki kekokohan (robustness) tinggi dan tidak rapuh terhadap deviasi minor penilaian pakar."
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_MUTED

    c2 = add_card(s10, Inches(0.8) + col_w + gap, Inches(1.7), col_w, Inches(5.0))
    tb2 = s10.shapes.add_textbox(Inches(1.0) + col_w + gap, Inches(1.9), col_w - Inches(0.4), Inches(4.6))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "2. Netralitas Geometris"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p.space_after = Pt(8)
    p = tf2.add_paragraph()
    p.text = "• Penggunaan Jarak Euclidean dalam TOPSIS memperlakukan seluruh dimensi secara ortogonal dan proporsional."
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_MAIN
    p.space_after = Pt(6)
    p = tf2.add_paragraph()
    p.text = "• Tidak terjadi fenomena rank reversal yang ekstrim ketika alternatif baru diujikan dalam skala simulasi."
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_MUTED

    c3 = add_card(s10, Inches(0.8) + (col_w + gap)*2, Inches(1.7), col_w, Inches(5.0))
    tb3 = s10.shapes.add_textbox(Inches(1.0) + (col_w + gap)*2, Inches(1.9), col_w - Inches(0.4), Inches(4.6))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "3. Validitas Domain Medis"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT
    p.space_after = Pt(8)
    p = tf3.add_paragraph()
    p.text = "• Bobot tertinggi AHP pada Prevalensi (37.46%) dan Volume Balita Stunting (24.24%) selaras dengan pedoman Kemenkes RI mengenai indikator intervensi spesifik."
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_MAIN
    p.space_after = Pt(6)
    p = tf3.add_paragraph()
    p.text = "• Kriteria sensitif (Sanitasi & Air Bersih) menjaga agar akar penyebab lingkungan tetap terakomodasi dalam skor akhir."
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_MUTED

    notes10 = s10.notes_slide.notes_text_frame
    notes10.text = (
        "SKRIP BICARA (ANALISIS SENSITIVITAS):\n"
        "Untuk memastikan model ini valid secara ilmiah, kami melakukan analisis sensitivitas. "
        "Hasil pengujian menunjukkan bahwa ketika bobot salah satu kriteria digeser naik-turun 15%, "
        "posisi klaster wilayah Sangat Tinggi (Brebes, Pemalang, Kebumen) tidak berubah secara drastis. "
        "Ini membuktikan bahwa model SPPK ini robust, stabil, dan tidak rentan terhadap bias penilai tunggal. "
        "Selain itu, perpaduan intervensi spesifik kesehatan dan intervensi sensitif lingkungan terakomodasi secara harmonis."
    )

    # =========================================================================
    # SLIDE 11: KESIMPULAN & REKOMENDASI KEBIJAKAN
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Penutup & Implikasi", "Kesimpulan & Rekomendasi Kebijakan", "Sintesis hasil penelitian dan kontribusi praktis bagi pembuat kebijakan")

    # 2 Big Cards: Kesimpulan vs Rekomendasi
    c_kes = add_card(s11, Inches(0.8), Inches(1.7), card_w2, Inches(5.0))
    tb_kes = s11.shapes.add_textbox(Inches(1.1), Inches(1.9), card_w2 - Inches(0.6), Inches(4.6))
    tfk = tb_kes.text_frame
    tfk.word_wrap = True
    p = tfk.paragraphs[0]
    p.text = "KESIMPULAN ILMIAH"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEAL_PRIMARY
    p.space_after = Pt(12)

    kes_items = [
        "1. Metode Hybrid AHP-TOPSIS terbukti berhasil menyelesaikan trade-off multi-kriteria untuk perankingan 35 Kabupaten/Kota secara efisien dan konsisten (CR = 0.0106).",
        "2. Algoritma TOPSIS secara akurat mengidentifikasi disparitas wilayah Jawa Tengah, dengan Kabupaten Brebes (0.9422) dan Pemalang (0.8055) sebagai wilayah paling mendesak untuk intervensi.",
        "3. Sistem web yang dikembangkan berhasil mewujudkan Explainable DSS dengan transparansi matematis di seluruh tahapan matriks perhitungan."
    ]
    for k in kes_items:
        p = tfk.add_paragraph()
        p.text = k
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(10)

    c_rek = add_card(s11, Inches(6.8), Inches(1.7), card_w2, Inches(5.0))
    tb_rek = s11.shapes.add_textbox(Inches(7.1), Inches(1.9), card_w2 - Inches(0.6), Inches(4.6))
    tfrk = tb_rek.text_frame
    tfrk.word_wrap = True
    p = tfrk.paragraphs[0]
    p.text = "REKOMENDASI KEBIJAKAN & PENGEMBANGAN"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = GOLD_ACCENT
    p.space_after = Pt(12)

    rek_items = [
        "1. Bagi Bappeda & Dinas Kesehatan Jateng: Rekomendasi perangkingan ini dapat dijadikan basis penentuan alokasi Bantuan Keuangan (Bankeu) penanganan kemiskinan ekstrem dan stunting.",
        "2. Intervensi Konvergen: Untuk Brebes dan Pemalang, fokus utama harus diarahkan pada dropping nakes dan pembangunan jamban/sanitasi sehat di kantong kemiskinan.",
        "3. Pengembangan Selanjutnya: Integrasi dengan Sistem Informasi Geografis (GIS/Peta Interaktif) dan modul prediksi runtun waktu (time-series forecasting) untuk tren tahun berikutnya."
    ]
    for r in rek_items:
        p = tfrk.add_paragraph()
        p.text = r
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MAIN
        p.space_after = Pt(10)

    notes11 = s11.notes_slide.notes_text_frame
    notes11.text = (
        "SKRIP BICARA (KESIMPULAN):\n"
        "Sebagai kesimpulan, penelitian ini membuktikan bahwa integrasi AHP dan TOPSIS mampu memecahkan permasalahan "
        "penentuan prioritas stunting secara akuntabel dan transparan. "
        "Hasil perangkingan ini dapat dimanfaatkan secara langsung oleh Bappeda dan Dinas Kesehatan Provinsi Jawa Tengah "
        "dalam mengunci alokasi anggaran intervensi gizi terpadu agar tepat sasaran dan terhindar dari pemborosan sumber daya. "
        "Untuk pengembangan masa depan, sistem dapat diperkaya dengan visualisasi spasial GIS."
    )

    # =========================================================================
    # SLIDE 12: ANTISIPASI PERTANYAAN DOSEN (Q&A DEFENSE PREPARATION)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Persiapan Ujian / Q&A", "Panduan Antisipasi Pertanyaan Kritis Dosen Penguji", "Strategi argumentasi akademik menjawab pertanyaan khas sidang SPPK / MCDM")

    # 4 Q&A Accordion Cards (2x2 Grid)
    qa_w = Inches(5.6)
    qa_h = Inches(2.3)

    qas = [
        ("Q1: Mengapa tidak memakai SAW atau WP?",
         "Jawaban: Simple Additive Weighting (SAW) dan Weighted Product (WP) mengabaikan konsep solusi ideal ganda (positif dan negatif) serta jarak euclidean. TOPSIS memiliki keunggulan geometris karena mempertimbangkan bukan hanya yang terbaik dari yang terbaik (A+), tetapi juga sejauh mungkin dari yang terburuk (A-).",
         TEAL_PRIMARY),
        ("Q2: Bagaimana jika nilai CR dari AHP melebihi 0.10?",
         "Jawaban: Jika CR > 0.10, sistem kami mendeteksi inkonsistensi secara otomatis. Pengguna/pakar wajib merevisi penilaian matriks perbandingannya hingga konsisten. Matriks default sistem kami memiliki CR = 0.0106 (1.06%), terbukti sangat konsisten.",
         CYAN_ACCENT),
        ("Q3: Mengapa Sanitasi & Air Bersih jadi kriteria 'Cost'?",
         "Jawaban: Dalam perspektif SPPK prioritas penanganan, wilayah dengan sanitasi dan air minum rendah adalah wilayah paling darurat yang menuntut intervensi segera. Karena itu nilai rendah harus menghasilkan skor preferensi penanganan yang tinggi.",
         GOLD_ACCENT),
        ("Q4: Bagaimana validitas data yang digunakan?",
         "Jawaban: Data indikator bersumber dari agregasi data BPS Jawa Tengah dan profil kesehatan Dinkes Jateng. Sistem juga mendukung fleksibilitas impor dataset multi-tahun (Excel) sehingga jika terdapat pembaharuan data resmi (SKDN / SSGI), sistem langsung mengkalkulasi ulang secara realtime.",
         GREEN_ACCENT),
    ]

    for idx, (qtitle, qans, qcol) in enumerate(qas):
        row_i = idx // 2
        col_i = idx % 2
        pos_x = Inches(0.8) + col_i * (qa_w + fgap_x)
        pos_y = Inches(1.7) + row_i * (qa_h + fgap_y)

        c = add_card(s12, pos_x, pos_y, qa_w, qa_h)
        tb = s12.shapes.add_textbox(pos_x + Inches(0.2), pos_y + Inches(0.12), qa_w - Inches(0.4), qa_h - Inches(0.25))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = qtitle
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = qcol
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = qans
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_MAIN

    notes12 = s12.notes_slide.notes_text_frame
    notes12.text = (
        "SKRIP BICARA (PENUTUP & Q&A):\n"
        "Demikian presentasi dari saya mengenai Sistem Pendukung Keputusan Penanganan Stunting di Jawa Tengah. "
        "Saya siap menerima kritik, saran, serta pertanyaan dari Bapak/Ibu dosen penguji untuk penyempurnaan sistem ini. "
        "Terima kasih atas perhatiannya. Waktu dan tempat saya kembalikan kepada ketua sidang."
    )

    output_path = r"D:\S2-UGM\SPPK\SPPK-Stunting\Presentasi_SPPK_Stunting_Jateng.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
